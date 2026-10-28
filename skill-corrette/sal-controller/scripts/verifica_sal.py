#!/usr/bin/env python3
"""Verifica i dati di input di un workbook SAL prima della consegna.

Uso:  python3 scripts/verifica_sal.py <workbook.xlsx> [--wp "WP1 FOH,WP2 CORRIDOI,..."]

Controlla i fogli "Libretto misure", "Contabilità SAL" e "Registro SAL" con la
struttura di references/formato-excel-sal.md. Non ricalcola le formule (openpyxl
non le esegue): controlla i dati da cui le formule dipendono.
Uscita 0 = nessun errore; 1 = errori trovati; 2 = file o fogli non leggibili.
Gli avvisi non bloccano ma vanno letti.
"""
import argparse
import sys
from datetime import date, datetime

try:
    import openpyxl
except ImportError:
    sys.exit("Manca la libreria openpyxl. Installala con: pip install openpyxl")

WP_DEFAULT = ["WP1 FOH", "WP2 CORRIDOI", "WP3 SUITE", "WP4 SCALA B", "WP5 20 ROOM"]
STATI_RILIEVO = ["montato", "parziale", "da verificare", "conteso"]
TIPI_VOCE = ["a corpo", "a misura"]
STATI_SAL = ["Emesso", "Verificato DLL", "Certificato", "Pagato", "Contestato"]

errori, avvisi = [], []


def norm(v):
    return str(v).strip() if v is not None else ""


def foglio(wb, nome):
    if nome not in wb.sheetnames:
        print(f"ERRORE: foglio '{nome}' non trovato. Fogli presenti: {', '.join(wb.sheetnames)}")
        sys.exit(2)
    return wb[nome]


def numero(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and v.startswith("="):
        return "formula"
    t = str(v).strip()
    if "," in t:  # formato italiano 1.234,50
        t = t.replace(".", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return "testo"


def leggi_voci(ws, wp_ammessi):
    voci = {}
    for r in range(2, ws.max_row + 1):
        cod, tipo = norm(ws[f"A{r}"].value), norm(ws[f"D{r}"].value).lower()
        if not cod or tipo not in TIPI_VOCE:
            if cod and tipo:
                errori.append(f"Contabilità SAL riga {r}: tipo '{tipo}' non valido per la voce {cod}. "
                              f"Valori ammessi: {', '.join(TIPI_VOCE)}")
            continue  # righe di subtotale o vuote
        wp = norm(ws[f"C{r}"].value)
        if wp not in wp_ammessi:
            errori.append(f"Contabilità SAL riga {r}: WP '{wp}' non riconosciuto per la voce {cod}. "
                          f"WP ammessi: {', '.join(wp_ammessi)}")
        if cod in voci:
            errori.append(f"Contabilità SAL riga {r}: voce {cod} duplicata (già alla riga {voci[cod]['riga']})")
        qta = numero(ws[f"F{r}"].value)
        if tipo == "a misura" and not isinstance(qta, float):
            errori.append(f"Contabilità SAL riga {r}: voce a misura {cod} senza quantità contrattuale numerica "
                          f"in colonna F (trovato: '{norm(ws[f'F{r}'].value)}')")
        voci[cod] = {"riga": r, "tipo": tipo, "um": norm(ws[f"E{r}"].value).lower(), "qta": qta, "wp": wp}
    if not voci:
        errori.append("Contabilità SAL: nessuna voce trovata. Servono codice in colonna A e tipo "
                      "'a corpo' o 'a misura' in colonna D, a partire dalla riga 2")
    return voci


def controlla_libretto(ws, voci, wp_ammessi):
    somme = {}
    ids = {}
    for r in range(2, ws.max_row + 1):
        rid = norm(ws[f"A{r}"].value)
        if not rid:
            continue
        if rid in ids:
            errori.append(f"Libretto misure riga {r}: ID rilievo {rid} duplicato (già alla riga {ids[rid]})")
        ids[rid] = r
        stato = norm(ws[f"M{r}"].value).lower()
        if stato not in STATI_RILIEVO:
            errori.append(f"Libretto misure riga {r} ({rid}): stato '{stato}' non valido. "
                          f"Valori ammessi: {', '.join(STATI_RILIEVO)}")
        wp = norm(ws[f"E{r}"].value)
        if wp not in wp_ammessi:
            errori.append(f"Libretto misure riga {r} ({rid}): WP '{wp}' non riconosciuto. "
                          f"WP ammessi: {', '.join(wp_ammessi)}")
        for col, campo in (("B", "data rilievo"), ("C", "rilevato da"), ("D", "SAL di competenza"),
                           ("F", "locale/area")):
            if not norm(ws[f"{col}{r}"].value):
                errori.append(f"Libretto misure riga {r} ({rid}): campo '{campo}' (colonna {col}) vuoto")
        if stato in ("montato", "parziale") and not norm(ws[f"L{r}"].value):
            errori.append(f"Libretto misure riga {r} ({rid}): stato '{stato}' senza riferimento documentale "
                          f"(colonna L vuota). Aggiungi foto/disegno/DDT o metti lo stato 'da verificare'")
        cod = norm(ws[f"G{r}"].value)
        if cod not in voci:
            errori.append(f"Libretto misure riga {r} ({rid}): codice voce '{cod}' non presente in "
                          f"Contabilità SAL (possibile extra/variante o errore di codifica)")
            continue
        v = voci[cod]
        um = norm(ws[f"I{r}"].value).lower()
        if um and v["um"] and um != v["um"]:
            errori.append(f"Libretto misure riga {r} ({rid}): UM '{um}' diversa da quella della voce "
                          f"{cod} ('{v['um']}')")
        qta, pct = numero(ws[f"J{r}"].value), numero(ws[f"K{r}"].value)
        if qta is not None and pct is not None:
            errori.append(f"Libretto misure riga {r} ({rid}): compilati sia quantità (J) sia % (K). "
                          f"Voce {cod} è {v['tipo']}: lascia solo "
                          f"{'la quantità' if v['tipo'] == 'a misura' else 'la %'}")
        elif v["tipo"] == "a misura" and pct is not None:
            errori.append(f"Libretto misure riga {r} ({rid}): voce a misura {cod} con % in colonna K; "
                          f"serve la quantità in colonna J")
        elif v["tipo"] == "a corpo" and qta is not None:
            errori.append(f"Libretto misure riga {r} ({rid}): voce a corpo {cod} con quantità in colonna J; "
                          f"serve la % di avanzamento in colonna K")
        if isinstance(pct, float) and not 0 <= pct <= 100:
            errori.append(f"Libretto misure riga {r} ({rid}): % avanzamento {pct} fuori da 0-100")
        if stato == "montato" and v["tipo"] == "a misura" and isinstance(qta, float):
            somme[cod] = somme.get(cod, 0) + qta
    for cod, tot in somme.items():
        q = voci[cod]["qta"]
        if isinstance(q, float) and tot > q + 1e-9:
            errori.append(f"Voce {cod}: quantità montata progressiva {tot:g} supera la contrattuale {q:g}. "
                          f"Serve una variante approvata (project-controller) o c'è un errore nei rilievi")
    if not ids:
        avvisi.append("Libretto misure: nessun rilievo trovato (colonna A vuota dalla riga 2)")


def controlla_registro(ws):
    prec_n, prec_data = None, None
    for r in range(2, ws.max_row + 1):
        n = norm(ws[f"A{r}"].value)
        if not n:
            continue
        stato = norm(ws[f"K{r}"].value)
        if stato and stato not in STATI_SAL:
            errori.append(f"Registro SAL riga {r} ({n}): stato '{stato}' non valido. "
                          f"Valori ammessi: {', '.join(STATI_SAL)}")
        cifre = "".join(c for c in n if c.isdigit())
        if cifre:
            num = int(cifre)
            if prec_n is not None and num != prec_n + 1:
                errori.append(f"Registro SAL riga {r}: {n} non segue il precedente (atteso n. {prec_n + 1})")
            prec_n = num
        d = ws[f"B{r}"].value
        if isinstance(d, (date, datetime)):
            if prec_data and d < prec_data:
                errori.append(f"Registro SAL riga {r} ({n}): data dati {d:%d/%m/%Y} precedente "
                              f"a quella del SAL prima ({prec_data:%d/%m/%Y})")
            prec_data = d
        else:
            avvisi.append(f"Registro SAL riga {r} ({n}): data dati in colonna B mancante o non in "
                          f"formato data (trovato: '{norm(d)}')")


def main():
    ap = argparse.ArgumentParser(description="Verifica i dati di input di un workbook SAL.")
    ap.add_argument("workbook")
    ap.add_argument("--wp", help="elenco WP ammessi separati da virgola (default: WP di 2514RW)")
    a = ap.parse_args()
    wp_ammessi = [w.strip() for w in a.wp.split(",")] if a.wp else WP_DEFAULT
    try:
        wb = openpyxl.load_workbook(a.workbook)
    except FileNotFoundError:
        print(f"ERRORE: file '{a.workbook}' non trovato")
        sys.exit(2)
    except Exception as e:
        print(f"ERRORE: il file '{a.workbook}' non è un .xlsx leggibile ({e})")
        sys.exit(2)
    voci = leggi_voci(foglio(wb, "Contabilità SAL"), wp_ammessi)
    controlla_libretto(foglio(wb, "Libretto misure"), voci, wp_ammessi)
    controlla_registro(foglio(wb, "Registro SAL"))
    for x in avvisi:
        print("AVVISO:", x)
    for x in errori:
        print("ERRORE:", x)
    print(f"\n{len(voci)} voci controllate — {len(errori)} errori, {len(avvisi)} avvisi.")
    sys.exit(1 if errori else 0)


if __name__ == "__main__":
    main()
